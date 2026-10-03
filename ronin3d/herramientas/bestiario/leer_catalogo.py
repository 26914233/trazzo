"""Paso 1 · Lee los dos catálogos xlsx del usuario y los vuelca a un JSON de trabajo.

Uso:
    python3 leer_catalogo.py <Bestiario_Universal_500.xlsx> <Bestiario_Universal_501-1000_Detallado.xlsx> <salida.json>

Solo se conservan las columnas con información propia de cada criatura (nombre, cultura,
región, categoría y, en el segundo archivo, la descripción). El resto de columnas del
catálogo original tienen el mismo texto de relleno en todas las filas.
Necesita openpyxl (pip install openpyxl).
"""
import json
import sys

import openpyxl


def filas(hoja, columnas):
    cabecera = [c.value for c in hoja[1]]
    posicion = {nombre: cabecera.index(nombre) for nombre in columnas.values()}
    salida = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if all(valor is None for valor in fila):
            continue
        salida.append({clave: fila[posicion[nombre]] for clave, nombre in columnas.items()})
    return salida


def main() -> None:
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    ruta_a, ruta_b, destino = sys.argv[1:]
    hoja_a = openpyxl.load_workbook(ruta_a, data_only=True)["Bestiario Universal"]
    hoja_b = openpyxl.load_workbook(ruta_b, data_only=True)["Bestiario 501-1000"]
    a = filas(hoja_a, {"id": "ID", "nombre": "Nombre", "cultura": "Cultura / Tradición",
                       "region": "Región", "categoria": "Categoría"})
    b = filas(hoja_b, {"id": "ID", "nombre": "Nombre", "cultura": "Tradición / Cultura",
                       "region": "Región", "categoria": "Categoría",
                       "descripcion": "Descripción detallada"})
    for fila in a:
        fila["archivo"] = "A"
        fila["descripcion"] = ""
    for fila in b:
        fila["archivo"] = "B"
    todo = a + b
    for fila in todo:
        fila["id"] = int(fila["id"])
    json.dump(todo, open(destino, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{len(a)} filas del archivo A + {len(b)} del archivo B = {len(todo)} en {destino}")


if __name__ == "__main__":
    main()
